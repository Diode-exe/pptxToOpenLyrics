import tkinter as tk
from tkinter import filedialog, messagebox

from pptx import Presentation
from create_ol import OpenLyricsSong

class GUI:
    def __init__(self):
        self.input_filename = None

        self.root = tk.Tk()
        self.root.title("PPTX to OpenLyrics Converter")
        self.root.bind("<Return>", lambda event: self.submit())

        self.title_var = tk.StringVar()
        self.title_entry_text = tk.Label(self.root, text="Title")
        self.title_entry_text.pack()
        self.title_entry = tk.Entry(self.root, textvariable=self.title_var)
        self.title_entry.pack()

        self.author_var = tk.StringVar()
        self.author_entry_text = tk.Label(self.root, text="Author")
        self.author_entry_text.pack()
        self.author_entry = tk.Entry(self.root, textvariable=self.author_var)
        self.author_entry.pack()


        self.file_chooser_button = tk.Button(self.root, text="Choose File", command=self.choose_file)
        self.file_chooser_button.pack()

        self.blank_slide_var = tk.BooleanVar()
        self.blank_slide_check = tk.Checkbutton(self.root, text="Blank slide at the end", variable=self.blank_slide_var)
        self.blank_slide_check.pack()

        self.submit_button = tk.Button(self.root, text="Submit", command=self.submit)
        self.submit_button.pack()

    def choose_file(self):
        self.input_filename = filedialog.askopenfilename(filetypes=[("PPTX files", "*.pptx")])

    def submit(self):
        if not self.input_filename:
            messagebox.showerror("No file selected", "Please choose a PPTX file first.")
            return

        self.title = self.title_var.get()
        self.author = self.author_var.get()
        self.blank_slide = self.blank_slide_var.get()

        converter = Converter()
        converter.convert(self.input_filename, self.title, self.author, self.blank_slide)

class Converter:
    def __init__(self):
        self.prs = None
        self.ol_song = OpenLyricsSong()

    def convert(self, input_filename, title, author, blank_slide):
        self.prs = Presentation(input_filename)
        root = self.ol_song.create_openlyrics(title, author)
        lyrics = root.find("lyrics")
        for i, slide in enumerate(self.prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    print(shape.text)
                    self.ol_song.add_verse(lyrics, f"verse{i+1}", [shape.text])
            if blank_slide:
                self.ol_song.add_verse(lyrics, f"verse{i+1}", [""])
        self.ol_song.pretty_and_save_xml(root, input_filename.replace(".pptx", ".xml"))

gui = GUI()
gui.root.mainloop()