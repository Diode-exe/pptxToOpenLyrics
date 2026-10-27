import time
import json
import os
from PySide6.QtCore import QFile, QObject, Signal # pylint: disable=E0611
from PySide6.QtWidgets import QApplication, QVBoxLayout, QProgressBar # pylint: disable=E0611
from PySide6.QtUiTools import QUiLoader # pylint: disable=E0611
from PySide6.QtWidgets import ( # pylint: disable=E0611
    QLabel,
    QPushButton,
    QFileDialog,
    QCheckBox,
    QMessageBox,
    QLineEdit,
)
import splitter

from pptx import Presentation
from create_ol import OpenLyricsSong

class GUI:
    def __init__(self):
        self.input_filename = None
        self.title = ""
        self.author = ""
        self.blank_slide = False
        self.file_picker_dir = ""
        self.dest_picker_dir = ""
        if not os.path.exists("src/settings.json"):
            self.create_settings()
        settings = self.load_settings_file()
        if settings:
            self.file_picker_dir = settings.get("file_picker_dir", "")
            self.dest_picker_dir = settings.get("dest_picker_dir", "")
        self.prs = None
        self.ol_song = None

        ui_file = QFile("src/main.ui")

        loader = QUiLoader()
        self.window = loader.load(ui_file)
        ui_file.close()

        if self.window is None:
            raise RuntimeError(f"Could not load src/main.ui: {loader.errorString()}")

        self.window.setWindowTitle("PPTX to OpenLyrics Converter")
        layout = QVBoxLayout(self.window.centralWidget())

        self.time_label = self.window.findChild(QLabel, "time_label")

        self.title_label = self.window.findChild(QLabel, "title_label")
        self.title_entry = self.window.findChild(QLineEdit, "title_entry")
        self.guess_title_button = self.window.findChild(QCheckBox, "guess_song_checkbox")
        # self.guess_title_button.stateChanged.connect(self.toggle_title_entry)

        self.author_label = self.window.findChild(QLabel, "author_label")
        self.author_entry = self.window.findChild(QLineEdit, "author_entry")

        self.file_chooser_button = self.window.findChild(QPushButton, "choose_file_button")
        self.file_chooser_button.clicked.connect(self.choose_file)

        self.dest_chooser_button = self.window.findChild(QPushButton, "choose_dest_button")
        self.dest_chooser_button.clicked.connect(self.choose_dest)

        self.blank_slide_var = self.window.findChild(QCheckBox, "blank_slide_checkbox")

        self.submit_button = self.window.findChild(QPushButton, "start_button")
        self.submit_button.clicked.connect(self.submit)

        self.progress_bar = self.window.findChild(QProgressBar, "progress_bar")
        self.progress_bar.setValue(0)
        
        self.splitter_button = self.window.findChild(QPushButton, "splitter_button")
        self.splitter_button.clicked.connect(self.show_splitter)

        self.window.show()

    def choose_file(self):
        self.input_filename = QFileDialog.getOpenFileName(
            self.window, "Select PPTX file", self.file_picker_dir, "PPTX files (*.pptx)",
        )[0]

    def choose_dest(self):
        self.dest_picker_dir = QFileDialog.getExistingDirectory(
            self.window, "Select Destination Directory", self.dest_picker_dir
        )

    def submit(self):
        if not self.input_filename:
            QMessageBox.critical(
                self.window, "No file selected", "Please choose a PPTX file first."
            )
            return

        self.title = self.title_entry.text()
        self.author = self.author_entry.text()
        self.blank_slide = self.blank_slide_var.isChecked()

        converter = Converter()
        converter.progress_signal.connect(self.progress_bar.setValue)
        conversion_time = converter.convert(self.input_filename, self.title, self.author, self.blank_slide, self.guess_title_button.isChecked(), self.dest_picker_dir)
        self.set_time(conversion_time)

        self.update_settings()

        QMessageBox.information(
            self.window,
            "Conversion Complete",
            "The PPTX file has been successfully converted to OpenLyrics XML.",
        )

    def toggle_title_entry(self, state):
        if state == 2:  # Checked
            self.title_entry.setDisabled(True)
        else:
            self.title_entry.setDisabled(False)

    def set_time(self, value):
        self.time_label.setText(f"Conversion took {value:.2f} seconds")

    def show_splitter(self):
        self.splitter_window = splitter.Splitter()

    def create_settings(self):
        settings = {
            "file_picker_dir": self.file_picker_dir,
            "dest_picker_dir": self.dest_picker_dir
        }
        with open("src/settings.json", "w", encoding="utf-8") as f:
            json.dump(settings, f)

    def load_settings_file(self):
        if not os.path.exists("src/settings.json"):
            return
        with open("src/settings.json", "r", encoding="utf-8") as f:
            settings = json.load(f)
        return settings

    def update_settings(self):
        if not os.path.exists("src/settings.json"):
            self.create_settings()
        with open("src/settings.json", "r", encoding="utf-8") as f:
            settings = json.load(f)
        settings["file_picker_dir"] = self.file_picker_dir
        settings["dest_picker_dir"] = self.dest_picker_dir
        with open("src/settings.json", "w", encoding="utf-8") as f:
            json.dump(settings, f)

class Converter(QObject):
    progress_signal = Signal(int) # pylint: disable=E0611

    def __init__(self):
        super().__init__()
        self.input_filename = None
        self.title = ""
        self.author = ""
        self.dest_picker_dir = ""
        self.blank_slide = False
        self.prs = None
        self.ol_song = None
        self.root = None
        self.lyrics = None

    def convert(self, input_filename, title, author, blank_slide, assume_title, dest_picker_dir):
        start_time = time.perf_counter()
        self.input_filename = input_filename
        self.title = title
        self.author = author
        self.blank_slide = blank_slide
        self.assume_title = assume_title
        self.dest_picker_dir = dest_picker_dir

        self.prs = Presentation(self.input_filename)
        self.ol_song = OpenLyricsSong()

        if self.assume_title:
            self.title = self.prs.slides[0].shapes[0].text if self.prs.slides and self.prs.slides[0].shapes or self.title else QMessageBox.critical(
                None, "No Title Found", "Please provide a title for the song."
            )

        self.root = self.ol_song.create_openlyrics(self.title, self.author)
        self.lyrics = self.root.find("lyrics")

        # Loop through slides and shape elements
        for i, slide in enumerate(self.prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    # for paragraph in shape.text_frame.paragraphs:
                        # if paragraph.text.strip():
                        #     print(paragraph.text)
                    print(shape.text)
                    self.ol_song.add_verse(self.lyrics, f"verse{i+1}", [shape.text])
            # self.progress_bar.setValue(int(((i + 1) / len(self.prs.slides)) * 100))
            self.progress_signal.emit(int(((i + 1) / len(self.prs.slides)) * 100))
            if self.blank_slide:
                self.ol_song.add_verse(self.lyrics, f"verse{i+1}", [""])

        self.ol_song.pretty_and_save_xml(self.root, os.path.join(self.dest_picker_dir, os.path.basename(self.input_filename.replace(".pptx", ".xml"))))
        end_time = time.perf_counter()
        print(f"Conversion time: {end_time - start_time:.2f} seconds")
        return end_time - start_time

if __name__ == "__main__":
    app = QApplication([])
    gui = GUI()
    app.exec()