import pptx
import PySide6.QtWidgets as QtWidgets
from PySide6.QtWidgets import QVBoxLayout
from PySide6.QtCore import QFile # pylint: disable=E0611
from PySide6.QtUiTools import QUiLoader # pylint: disable=E0611

class Splitter:
    def __init__(self):
        self.prs = None
        ui_file = QFile("src/splitter.ui")

        loader = QUiLoader()
        self.window = loader.load(ui_file)
        ui_file.close()
        self.window.setWindowTitle("Slide Splitter")

        self.file_chooser_button = self.window.findChild(QtWidgets.QPushButton, "choose_file_button")
        self.file_chooser_button.clicked.connect(self.choose_file)

        self.checkbox_scroll_area_contents = self.window.findChild(QtWidgets.QWidget, "checkbox_scroll_area_contents")
        if not self.checkbox_scroll_area_contents:
            raise ValueError("Could not find checkbox_scroll_area_contents in the UI")

        self.window.show()

    def choose_file(self):
        filename, _ = QtWidgets.QFileDialog.getOpenFileName(
            self.window, "Select PPTX file", "", "PPTX files (*.pptx)"
        )
        if filename:
            self.prs = pptx.Presentation(filename)
            self.render_checkboxes()

    def get_slides(self):
        return self.prs.slides

    def get_slide_count(self):
        return len(self.prs.slides)

    def render_checkboxes(self):
        checkbox_src = []
        checkboxes = []
        scroll_layout = QVBoxLayout(self.checkbox_scroll_area_contents)
        for i, _ in enumerate(self.get_slides()):
            checkbox_src.append(f"Slide {i + 1}")
            print(f"Adding checkbox for Slide {i + 1}")
        for checkbox in checkbox_src:
            checkbox_widget = QtWidgets.QCheckBox(checkbox)
            checkboxes.append(checkbox_widget)
            scroll_layout.addWidget(checkbox_widget)
        return checkboxes